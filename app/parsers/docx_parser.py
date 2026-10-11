"""DOCX parser using python-docx.

We map Word heading styles to section-aware blocks so structure-aware chunking and section
citations work. Paragraphs under a heading inherit that heading as their section.
"""

from __future__ import annotations

import io

from app.domain import DocumentBlock, DocumentMetadata, ParsedDocument
from app.parsers.base import ParserError, blocks_to_text, normalize_text


class DocxParser:
    extensions = (".docx",)

    def parse(self, data: bytes, metadata: DocumentMetadata) -> ParsedDocument:
        try:
            import docx  # python-docx
        except ImportError as exc:  # pragma: no cover
            raise ParserError("python-docx is not installed") from exc

        try:
            document = docx.Document(io.BytesIO(data))
        except Exception as exc:
            raise ParserError(f"Invalid or corrupted DOCX: {exc}") from exc

        blocks: list[DocumentBlock] = []
        current_section: str | None = None

        for para in document.paragraphs:
            text = normalize_text(para.text)
            if not text:
                continue
            style = (para.style.name if para.style else "") or ""
            if style.lower().startswith("heading") or style.lower() == "title":
                level = _heading_level(style)
                current_section = text
                blocks.append(DocumentBlock(text=text, kind="heading", section=text, level=level))
                if metadata.title is None and style.lower() == "title":
                    metadata.title = text
            else:
                blocks.append(DocumentBlock(text=text, kind="paragraph", section=current_section))

        if not blocks:
            raise ParserError("DOCX contained no extractable text")
        return ParsedDocument(metadata=metadata, text=blocks_to_text(blocks), blocks=blocks)


def _heading_level(style_name: str) -> int:
    digits = "".join(c for c in style_name if c.isdigit())
    return int(digits) if digits else 1
